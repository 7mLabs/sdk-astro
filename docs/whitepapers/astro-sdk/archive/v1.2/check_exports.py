from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import json,re,hashlib
P=Path(__file__).resolve().parent.parent
d=Document(P/'Paper Astro SDK 7MLabs.docx');pdf=PdfReader(P/'Paper Astro SDK 7MLabs.pdf').pages
parts=[]
for el in d.element.body:
 if el.tag==qn('w:p'):parts.append(''.join(el.itertext()) if False else ''.join(t.text or '' for t in el.iter(qn('w:t'))))
 elif el.tag==qn('w:tbl'):
  for row in el.findall(qn('w:tr')):
   for cell in row.findall(qn('w:tc')):parts.append(''.join(t.text or '' for t in cell.iter(qn('w:t'))))
actual=[]
repeated_headers=[]
for page_number,page in enumerate(pdf):
 lines=page.extract_text().splitlines()
 if lines and lines[-1].strip()==str(page_number+1):lines.pop()
 if lines and lines[0].strip()==str(page_number+1):lines.pop(0)
 # Known repeated header, visually verified on page 7 in v1.2.
 if page_number == 6 and lines and re.sub(r'\s+','',lines[0]) == 'MẫuNguồnbyteContextbyteR%':
  repeated_headers.append({'page':7,'text':lines.pop(0)})
 actual.append('\n'.join(lines))
def norm(x):return re.sub(r'\s+','',x)
a=norm('\n'.join(parts));b=norm('\n'.join(actual))
result={'repeatedTableHeadersExcluded':repeated_headers,'pages':len(pdf),'editableParagraphs':len(d.paragraphs),'editableTables':len(d.tables),'orderedTextParity':a==b,'docxCharsWithoutWhitespace':len(a),'pdfCharsWithoutWhitespace':len(b),'files':{f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in P.glob('Paper*')}}
if a!=b:
 for i,(x,y) in enumerate(zip(a,b)):
  if x!=y:result['firstDifference']=[i,a[max(0,i-70):i+100],b[max(0,i-70):i+100]];break
(P/'evidence/export-check.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False,indent=2))
assert a==b
