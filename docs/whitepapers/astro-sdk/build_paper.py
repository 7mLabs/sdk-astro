from pathlib import Path
import json,re
from equations import word_equation
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).parent
r=json.loads((P/'evidence/review-v1.1/results.json').read_text())
def num(n): return f'{n:,}'.replace(',','.')
rows=[['Mẫu','Nguồn byte','Context byte','R %']]+[[x['fixture'].replace('-result.json',''),num(x['inputBytes']),num(x['outputBytes']),f"{x['reductionPercent']:.2f}".replace('.',',')] for x in r['compression']]
refs=[['Trường hợp','Swiss','Skyfield']]+[[x['case'],format(x['swiss']['maxLongitudeErrorDegrees'],'.10f').rstrip('0').rstrip('.').replace('.',','),f"{x['skyfield']['maxLongitudeErrorDegrees']:.10f}".replace('.',',')] for x in r['references']]
md=(P/'paper-source.md').read_text().replace('{{PYTHON_VERSION}}',r['python'])
for name,table in [('COMPRESSION_TABLE',rows),('REFERENCE_TABLE',refs)]:
 text='\n'.join(['| '+' | '.join(table[0])+' |','| '+' | '.join(['---']*len(table[0]))+' |']+['| '+' | '.join(row)+' |' for row in table[1:]])
 md=md.replace('{{'+name+'}}',text)
(P/'whitepaper.md').write_text(md)
d=Document(); sec=d.sections[0]
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.footer_distance=Inches(.5)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 s=d.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
 s.font.size=Pt(11);s.paragraph_format.line_spacing=1
normal=d.styles['Normal'];normal.paragraph_format.first_line_indent=Pt(16)
normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.widow_control=True
for name,size,before,after in [('Title',14,0,10),('Heading 1',12,12,5),('Heading 2',11,9,4)]:
 s=d.styles[name];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after);s.paragraph_format.first_line_indent=Pt(0);s.paragraph_format.keep_with_next=True
cap=d.styles['Caption'];cap.font.size=Pt(10);cap.font.italic=True;cap.font.bold=False;cap.paragraph_format.first_line_indent=Pt(0);cap.paragraph_format.space_before=Pt(7);cap.paragraph_format.space_after=Pt(4);cap.paragraph_format.keep_with_next=True
for name in ['First','Reference','Meta']:
 s=d.styles.add_style(name,1);s.base_style=normal
 s.paragraph_format.first_line_indent=Pt(0)
d.styles['Reference'].font.size=Pt(10)
d.styles['Reference'].paragraph_format.left_indent=Pt(18);d.styles['Reference'].paragraph_format.first_line_indent=Pt(-18);d.styles['Reference'].paragraph_format.space_after=Pt(4);d.styles['Reference'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
d.styles['Meta'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;d.styles['Meta'].font.size=Pt(10);d.styles['Meta'].paragraph_format.space_after=Pt(8)
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
def table(lines):
 data=[[c.strip() for c in ln.strip('|').split('|')] for ln in lines if not re.match(r'^\|[\s|:-]+\|$',ln)]
 t=d.add_table(rows=0,cols=len(data[0]));t.autofit=False
 numeric=data[0][0] in ['Mẫu','Trường hợp','Nhóm']
 if len(data[0])==4: widths=[2.43,1.28,1.28,1.28]
 elif len(data[0])==3: widths=[3.05,1.45,1.77] if numeric else [1.32,2.45,2.50]
 else: widths=[4.9,1.37] if numeric else [1.85,4.42]
 for col,w in zip(t.columns,widths): col.width=Inches(w)
 for i,row in enumerate(data):
  cells=t.add_row().cells
  for j,(cell,txt) in enumerate(zip(cells,row)):
   cell.width=Inches(widths[j]);p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or not numeric else WD_ALIGN_PARAGRAPH.RIGHT
   run=p.add_run(txt);run.font.size=Pt(10);run.bold=i==0
  pr=t.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0: pr.append(OxmlElement('w:tblHeader'))
 props=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
 for edge in ['top','bottom']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');b.append(e)
 props.append(b)
def linked_text(p,text):
 text=re.sub(r'([BCRw])_([ij])',lambda m:m[1]+{'i':'ᵢ','j':'ⱼ'}[m[2]],text)
 from docx.opc.constants import RELATIONSHIP_TYPE as RT
 last=0
 for m in re.finditer(r'https?://[^\s]+',text):
  p.add_run(text[last:m.start()])
  href=m.group();h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(href,RT.HYPERLINK,is_external=True))
  run=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'000000');pr.append(c);run.append(pr)
  t=OxmlElement('w:t');t.text=href;run.append(t);h.append(run);p._p.append(h);last=m.end()
 p.add_run(text[last:])
first=True;inrefs=False;lines=md.splitlines();i=0
while i<len(lines):
 ln=lines[i].strip()
 if not ln: i+=1;continue
 if ln.startswith('|'):
  block=[]
  while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
  table(block);first=False;continue
 if ln.startswith('{{EQ_'):
  key=ln[5:-2];p=d.add_paragraph(style='First');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(8)
  word_equation(key,p)
 elif ln.startswith('# '):
  p=d.add_paragraph(ln[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 elif ln.startswith('### '): d.add_paragraph(ln[4:],'Heading 2');first=True
 elif ln.startswith('## '):
  title=ln[3:];d.add_paragraph(title,'Heading 1');first=True;inrefs=title=='Tài liệu tham khảo'
 elif ln.startswith('7MLabs ·'):d.add_paragraph(ln,'Meta')
 elif ln.startswith('!['):
  match=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',ln)
  p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(6)
  picture=p.add_run().add_picture(str(P/match.group(2)),width=Inches(6.26))
  picture._inline.docPr.set('descr',match.group(1))
 elif re.match(r'^Hình \d+\.',ln):
  p=d.add_paragraph(ln,'Caption');p.paragraph_format.keep_with_next=False
 elif ln.startswith('Công thức ('):
  p=d.add_paragraph(ln,'First');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6)
 elif re.match(r'^Bảng \d+\.',ln):d.add_paragraph(ln,'Caption')
 else:
  style='Reference' if inrefs else ('First' if first or ln.startswith('Từ khóa') else 'Normal')
  p=d.add_paragraph(style=style);linked_text(p,ln);first=False
 i+=1
for style in d.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
 if style.type == 1:
  rf=style.element.get_or_add_rPr().find(qn('w:rFonts'))
  if rf is not None:
   for key in list(rf.attrib):
    if key.endswith('Theme'): del rf.attrib[key]
   for key in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+key),'Times New Roman')
d.core_properties.title='Astro SDK và khả năng bảo toàn dữ liệu thiên văn'
d.core_properties.author='7MLabs';d.core_properties.subject='Nghiên cứu trường hợp phần mềm';d.core_properties.version='1.4'
d.save(P/'Paper Astro SDK 7MLabs.docx')
print('Created DOCX; tests',r['passedTests'],'reduction',100*(1-r['totalOutputBytes']/r['totalInputBytes']))
